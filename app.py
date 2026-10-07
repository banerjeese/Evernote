import io, os, re, html
from flask import Flask, Response, jsonify, request, send_file, send_from_directory
from mysql.connector import pooling

DB = dict(
    host=os.getenv("DB_HOST", "100.125.233.115"),
    port=int(os.getenv("DB_PORT", "3306")),
    user=os.getenv("DB_USER", "dbuser"),
    password=os.getenv("DB_PASSWORD", "232527"),
    database=os.getenv("DB_NAME", "ProductivityAnalytics"),
    charset="utf8mb4",
)
USER_ID = 1  # single user, no auth

pool = pooling.MySQLConnectionPool(
    pool_name="notes", pool_size=8, pool_reset_session=False, autocommit=True, **DB
)
app = Flask(__name__, static_folder="static", static_url_path="")
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024  # max upload (images)


def snip(raw):
    t = (raw or "").replace("<!--rt-->", "")
    t = re.sub(r'<div class="i"[^>]*></div>', " 🖼 ", t)
    t = re.sub(r"<[^>]*$", "", t)
    t = html.unescape(re.sub(r"<[^>]+>", " ", t))
    return re.sub(r"\s+", " ", t).strip()[:140]


def run(sql, args=(), fetch="all"):
    """fetch: 'all' | 'one' | 'write' (returns (lastrowid, rowcount))."""
    conn = pool.get_connection()
    try:
        cur = conn.cursor(dictionary=True)
        cur.execute(sql, args)
        if fetch == "write":
            out = (cur.lastrowid, cur.rowcount)
        elif fetch == "one":
            out = cur.fetchone()
        else:
            out = cur.fetchall()
        cur.close()
        return out
    finally:
        conn.close()


D = "CAST({0} AS CHAR)"            # 2026-10-04
T = "LEFT(CAST({0} AS CHAR),5)"    # 10:30


IMG_SQL = """CREATE TABLE IF NOT EXISTS note_images (
  image_id   INT AUTO_INCREMENT PRIMARY KEY,
  subnote_id INT NOT NULL,
  mime       VARCHAR(40) NOT NULL,
  data       MEDIUMBLOB NOT NULL,
  created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
  KEY idx_subnote (subnote_id)
) ENGINE=InnoDB"""
try:
    run(IMG_SQL, (), "write")
    IMG_OK = True
except Exception as e:  # e.g. dbuser has no CREATE permission
    IMG_OK = False
    print("!! Could not create table note_images:", e)
    print("!! Run this once in MySQL, then restart:\n" + IMG_SQL)


@app.get("/")
def index():
    return send_from_directory("static", "index.html")


# ---------- notes ----------
@app.get("/api/notes")
def list_notes():
    return jsonify(run(
        f"""SELECT n.note_id id, n.note_topic topic, {D.format('n.note_date')} date,
                   {T.format('n.note_time')} time,
                   (SELECT COUNT(*) FROM subnotes s WHERE s.note_id = n.note_id) cnt
            FROM notes n WHERE n.userId=%s ORDER BY n.note_id DESC""", (USER_ID,)))

@app.get("/api/pages")
def list_pages():
    return jsonify(run(
        """SELECT n.note_id, n.note_topic,
                  s.subnote_id, s.subnote_topic
           FROM notes n
           LEFT JOIN subnotes s ON s.note_id = n.note_id
           WHERE n.userId = %s
           ORDER BY n.note_id DESC, s.subnote_id DESC""",
        (USER_ID,)))

@app.post("/api/notes")
def add_note():
    topic = (request.json.get("topic") or "Untitled").strip()[:255]
    nid, _ = run("INSERT INTO notes (userId, note_topic, note_date, note_time) "
                 "VALUES (%s,%s,CURDATE(),CURTIME())", (USER_ID, topic), "write")
    return jsonify(run(
        f"""SELECT note_id id, note_topic topic, {D.format('note_date')} date,
                   {T.format('note_time')} time, 0 cnt FROM notes WHERE note_id=%s""",
        (nid,), "one")), 201


@app.patch("/api/notes/<int:nid>")
def rename_note(nid):
    topic = (request.json.get("topic") or "Untitled").strip()[:255]
    run("UPDATE notes SET note_topic=%s WHERE note_id=%s AND userId=%s",
        (topic, nid, USER_ID), "write")
    return jsonify(ok=True)


@app.delete("/api/notes/<int:nid>")
def del_note(nid):
    if IMG_OK:
        run("DELETE i FROM note_images i JOIN subnotes s ON s.subnote_id=i.subnote_id "
            "JOIN notes n ON n.note_id=s.note_id WHERE s.note_id=%s AND n.userId=%s",
            (nid, USER_ID), "write")
    run("DELETE s FROM subnotes s JOIN notes n ON n.note_id=s.note_id "
        "WHERE s.note_id=%s AND n.userId=%s", (nid, USER_ID), "write")
    run("DELETE FROM notes WHERE note_id=%s AND userId=%s", (nid, USER_ID), "write")
    return jsonify(ok=True)


# ---------- subnotes ----------
@app.get("/api/notes/<int:nid>/subnotes")
def list_subs(nid):
    # snippet only: full description is fetched when a subnote is opened
    rows = run(
        f"""SELECT s.subnote_id id, s.subnote_topic topic,
                   LEFT(IFNULL(s.description,''),800) snippet,
                   {D.format('s.subnote_date')} date, {T.format('s.subnote_time')} time
            FROM subnotes s JOIN notes n ON n.note_id=s.note_id
            WHERE s.note_id=%s AND n.userId=%s ORDER BY s.subnote_id DESC""",
        (nid, USER_ID))
    for r in rows:
        r["snippet"] = snip(r["snippet"])
    return jsonify(rows)


@app.post("/api/notes/<int:nid>/subnotes")
def add_sub(nid):
    j = request.json or {}
    topic = (j.get("topic") or "Untitled").strip()[:255]
    sid, _ = run(
        "INSERT INTO subnotes (note_id, subnote_topic, description, subnote_date, subnote_time) "
        "SELECT note_id,%s,%s,CURDATE(),CURTIME() FROM notes WHERE note_id=%s AND userId=%s",
        (topic, j.get("description", ""), nid, USER_ID), "write")
    return jsonify(run(
        f"""SELECT subnote_id id, subnote_topic topic, '' snippet,
                   {D.format('subnote_date')} date, {T.format('subnote_time')} time
            FROM subnotes WHERE subnote_id=%s""", (sid,), "one")), 201


@app.get("/api/subnotes/<int:sid>")
def get_sub(sid):
    row = run("SELECT s.subnote_id id, s.subnote_topic topic, IFNULL(s.description,'') description "
              "FROM subnotes s JOIN notes n ON n.note_id=s.note_id "
              "WHERE s.subnote_id=%s AND n.userId=%s", (sid, USER_ID), "one")
    return (jsonify(row), 200) if row else (jsonify(error="not found"), 404)


@app.patch("/api/subnotes/<int:sid>")
def edit_sub(sid):
    j = request.json or {}
    sets, args = [], []
    if "topic" in j:
        sets.append("s.subnote_topic=%s"); args.append((j["topic"] or "Untitled").strip()[:255])
    if "description" in j:
        sets.append("s.description=%s"); args.append(j["description"])
    if sets:
        run(f"UPDATE subnotes s JOIN notes n ON n.note_id=s.note_id SET {','.join(sets)} "
            "WHERE s.subnote_id=%s AND n.userId=%s", (*args, sid, USER_ID), "write")
    return jsonify(ok=True)


@app.delete("/api/subnotes/<int:sid>")
def del_sub(sid):
    if IMG_OK:
        run("DELETE i FROM note_images i JOIN subnotes s ON s.subnote_id=i.subnote_id "
            "JOIN notes n ON n.note_id=s.note_id WHERE i.subnote_id=%s AND n.userId=%s",
            (sid, USER_ID), "write")
    run("DELETE s FROM subnotes s JOIN notes n ON n.note_id=s.note_id "
        "WHERE s.subnote_id=%s AND n.userId=%s", (sid, USER_ID), "write")
    return jsonify(ok=True)


# ---------- export (docx) ----------
def parse_blocks(raw):
    """Saved description -> list of lines: {t: p|b|h|img, text, l(indent), id, w}"""
    raw = raw or ""
    if not raw.startswith("<!--rt-->"):  # old plain-text notes
        return [dict(t="p", text=ln, l=0) for ln in raw.split("\n")]
    out = []
    for m in re.finditer(r"<div([^>]*)>(.*?)</div>", raw[9:], re.S):
        attrs, inner = m.groups()
        cls = (re.search(r'class="([^"]*)"', attrs) or [None, ""])[1]
        lvl = int((re.search(r'data-l="(\d)"', attrs) or [None, 0])[1])
        if cls.split(" ")[0] == "i":
            idm, wm = re.search(r'data-id="(\d+)"', attrs), re.search(r'data-w="(\d+)"', attrs)
            if idm:
                out.append(dict(t="img", id=int(idm.group(1)), w=int(wm.group(1)) if wm else 0, l=0))
            continue
        text = html.unescape(re.sub(r"<[^>]+>", "", inner)).replace("\u00a0", " ")
        out.append(dict(t={"b": "b", "h": "h", "t": "t"}.get(cls, "p"), text=text, l=lvl, c='data-c="1"' in attrs))
    return out


def load_sub(sid):
    return run(
        f"""SELECT s.subnote_topic topic, IFNULL(s.description,'') description,
                   {D.format('s.subnote_date')} date, {T.format('s.subnote_time')} time
            FROM subnotes s JOIN notes n ON n.note_id=s.note_id
            WHERE s.subnote_id=%s AND n.userId=%s""", (sid, USER_ID), "one")


def split_title(t):
    """'Topic {by: Name}' -> ('Topic', 'by: Name')"""
    t = (t or "").strip()
    m = re.match(r"^(.*?)\s*\{\s*(by\b[^}]*?)\s*\}\s*$", t, re.I)
    return (m.group(1).strip(), m.group(2).strip()) if m and m.group(1).strip() else (t or "Untitled", "")


def build_docx(row, imgs):
    from docx import Document
    from docx.image.image import Image as DImage
    from docx.shared import Emu, Inches, Pt, RGBColor
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    title, by = split_title(row["topic"])
    r = d.add_paragraph().add_run("✎ AdiNotes")
    r.font.size, r.font.color.rgb = Pt(9), RGBColor(0x88, 0x88, 0x88)
    d.add_heading(title, level=1)
    if by:
        r = d.add_paragraph().add_run(by)
        r.font.size, r.font.color.rgb = Pt(11), RGBColor(0x88, 0x88, 0x88)
    sec = d.sections[0]
    avail = int(sec.page_width - sec.left_margin - sec.right_margin)
    for b in parse_blocks(row["description"]):
        if b["t"] == "img":
            try:
                data = imgs[b["id"]][1]
                native = int(DImage.from_blob(data).px_width / 96 * 914400)
                width = int(avail * b["w"] / 100) if b["w"] else min(native, avail)
                d.add_paragraph().add_run().add_picture(io.BytesIO(data), width=Emu(width))
            except Exception:
                d.add_paragraph("[image could not be added]")
        elif b["t"] == "h":
            d.add_heading(b["text"], level=2).paragraph_format.left_indent = Inches(0.35 * b["l"])
        else:
            p = d.add_paragraph()
            pf = p.paragraph_format
            pf.space_after = Pt(2)
            if b["t"] == "t":                      # checkbox line
                pf.left_indent, pf.first_line_indent = Inches(0.35 * b["l"] + 0.3), Inches(-0.3)
                p.add_run("☑ " if b.get("c") else "☐ ")
                tr = p.add_run(b["text"])
                if b.get("c"):
                    tr.font.strike, tr.font.color.rgb = True, RGBColor(0x88, 0x88, 0x88)
            elif b["t"] == "b":
                pf.left_indent, pf.first_line_indent = Inches(0.35 * b["l"] + 0.25), Inches(-0.2)
                p.add_run("• " + b["text"])
            else:
                pf.left_indent = Inches(0.35 * b["l"])
                p.add_run(b["text"])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@app.get("/api/subnotes/<int:sid>/export.<fmt>")
def export_sub(sid, fmt):
    if fmt != "docx":
        return "", 404
    row = load_sub(sid)
    if not row:
        return "", 404
    try:
        import docx  # noqa: F401
    except ImportError:
        return "Please run:  pip install python-docx   (then restart the app)", 501
    imgs = {}
    if IMG_OK:
        for r in run("SELECT image_id, mime, data FROM note_images WHERE subnote_id=%s", (sid,)):
            imgs[r["image_id"]] = (r["mime"], bytes(r["data"]))
    name = (re.sub(r'[\\/:*?"<>|\r\n]+', "-", split_title(row["topic"])[0]).strip()[:80] or "note") + ".docx"
    resp = send_file(io.BytesIO(build_docx(row, imgs)),
                     mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                     as_attachment=True, download_name=name)
    resp.headers["Cache-Control"] = "no-store"
    return resp


# ---------- images (stored in MySQL, served with long cache) ----------
ALLOWED = {"image/png", "image/jpeg", "image/gif", "image/webp"}


@app.post("/api/subnotes/<int:sid>/images")
def add_image(sid):
    if not IMG_OK:
        return jsonify(error="note_images table is missing"), 503
    mime, data = (request.mimetype or "").lower(), request.get_data()
    if mime not in ALLOWED or not data:
        return jsonify(error="unsupported image"), 400
    iid, n = run("INSERT INTO note_images (subnote_id, mime, data) "
                 "SELECT s.subnote_id,%s,%s FROM subnotes s JOIN notes n ON n.note_id=s.note_id "
                 "WHERE s.subnote_id=%s AND n.userId=%s", (mime, data, sid, USER_ID), "write")
    return (jsonify(id=iid), 201) if n else (jsonify(error="not found"), 404)


@app.get("/api/images/<int:iid>")
def get_image(iid):
    row = IMG_OK and run(
        "SELECT i.mime, i.data FROM note_images i JOIN subnotes s ON s.subnote_id=i.subnote_id "
        "JOIN notes n ON n.note_id=s.note_id WHERE i.image_id=%s AND n.userId=%s",
        (iid, USER_ID), "one")
    if not row:
        return "", 404
    r = Response(bytes(row["data"]), mimetype=row["mime"])
    r.headers["Cache-Control"] = "private, max-age=31536000, immutable"
    return r


@app.delete("/api/images/<int:iid>")
def del_image(iid):
    if IMG_OK:
        run("DELETE i FROM note_images i JOIN subnotes s ON s.subnote_id=i.subnote_id "
            "JOIN notes n ON n.note_id=s.note_id WHERE i.image_id=%s AND n.userId=%s",
            (iid, USER_ID), "write")
    return jsonify(ok=True)


if __name__ == "__main__":
    from waitress import serve
    print("Running on http://localhost:4500")
    serve(app, host="0.0.0.0", port=4500, threads=8)